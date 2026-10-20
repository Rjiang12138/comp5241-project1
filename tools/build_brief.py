"""Build the 2-page project brief (DOCX) from the prototype screenshots."""

from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SHOTS = ROOT / "docs" / "screenshots"
PRINT = SHOTS / "print"
OUT = ROOT / "docs" / "Praxis-project-brief.docx"

INK = RGBColor(0x10, 0x18, 0x28)
BODY = RGBColor(0x24, 0x2B, 0x3A)
MUTED = RGBColor(0x66, 0x70, 0x85)


def set_border_bottom(paragraph, color="D9DDE3", size=6):
    p = paragraph._p
    pPr = p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def para(doc, text, size=10.5, color=BODY, bold=False, before=0, after=7,
         spacing=1.14, align=WD_ALIGN_PARAGRAPH.LEFT, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = spacing
    p.alignment = align
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.bold = bold
    run.italic = italic
    return p


def heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(15)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = INK
    return p


def caption(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run(text)
    run.font.size = Pt(8)
    run.font.color.rgb = MUTED
    return p


def picture(doc, name, width_cm, cap):
    doc.add_picture(str(print_copy(name)), width=Cm(width_cm))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.paragraphs[-1].paragraph_format.space_after = Pt(0)
    caption(doc, cap)


def print_copy(name: str, px: int = 1700) -> Path:
    """JPEG copy of a screenshot at print resolution (LibreOffice drops the PNGs)."""
    PRINT.mkdir(parents=True, exist_ok=True)
    out = PRINT / (Path(name).stem + ".jpg")
    if not out.exists():
        img = Image.open(SHOTS / name).convert("RGB")
        if img.width > px:
            img = img.resize((px, round(img.height * px / img.width)), Image.LANCZOS)
        img.save(out, quality=92, optimize=True)
    return out



def remove_table_borders(table):
    tblPr = table._tbl.tblPr
    for existing in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(existing)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    # keep schema order: tblBorders must precede shd / tblLayout / tblCellMar / tblLook
    tblPr.insert_element_before(
        borders,
        "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook",
        "w:tblCaption", "w:tblDescription",
    )


def build():
    doc = Document()

    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.top_margin = Cm(1.9)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(2.1)
    section.right_margin = Cm(2.1)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)

    # ---------------------------------------------------------------- title
    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(2)
    run = title.add_run("Praxis")
    run.bold = True
    run.font.size = Pt(24)
    run.font.color.rgb = INK

    para(doc, "A GenAI course studio for software engineering education",
         size=12, color=MUTED, after=4)

    meta = para(doc, "COMP5241 Software Engineering  ·  Group project  ·  20 October 2026  ·  "
                     "github.com/Rjiang12138/comp5241-project1",
                size=8.5, color=MUTED, after=2)
    set_border_bottom(meta)

    # ---------------------------------------------------------------- body
    heading(doc, "Why we built it")
    para(doc, "Students meet software engineering as a list of definitions. In the same week they meet a "
              "coding agent that can produce a working app in an afternoon, and the two things do not "
              "connect. Praxis is our attempt to connect them: a studio where an instructor turns a topic "
              "or their own lecture deck into a short course that makes students practise the part the "
              "agent cannot do for them. Framing a requirement, reviewing a diff, defending a design "
              "decision.")

    heading(doc, "What it does")
    para(doc, "An instructor writes a brief, picks an audience and a time budget, and gets a five-module "
              "draft. Each module is assigned an activity format: reading, stakeholder simulation, code "
              "lab, design-review role play, quiz. The instructor edits, reorders or deletes anything. A "
              "panel on the right records what the generator did, step by step, so the draft can be "
              "argued with instead of accepted.")

    picture(doc, "studio.png", 16.3,
            "The studio. Brief and audience on the left, an editable draft in the middle, the generator's "
            "log and a pedagogy check on the right.")

    heading(doc, "What is different")
    para(doc, "1.  Three registers, one source. The same module is written three ways: an analogy-first "
              "version for a business student, a version built on contracts and invariants for a CS "
              "graduate, and a formal one. It is a switch inside the lesson, not three separate courses.")
    para(doc, "2.  Paths from a goal. A learner says what they want to be able to do and how many hours a "
              "week they have. The platform sequences what already exists, generates what does not, and "
              "gives a reason for each step.")
    para(doc, "3.  Generation you can audit. Five logged steps, the sources it read, and a comparison with "
              "the previous draft. The log exists because a draft that arrives complete is very hard to "
              "argue with.")

    heading(doc, "Where we are")
    para(doc, "Six screens, clickable, plain HTML, CSS and JavaScript with no build step. It runs from a "
              "file on any laptop, which is how we demo it. Students work through a module inside a player; "
              "the code lab starts with a failing test, a hint that only appears after one failed attempt, "
              "and a tutor that answers from the course material and says which lecture it came from.")
    para(doc, "Nothing is deployed yet; the backend is weeks 9 to 11. We are not building the mini-game or "
              "the AI-narrated slides this time. Both were on the format menu and both cost more to make "
              "convincing than they return, so we would rather finish five activity formats properly.")

    # two screenshots inline, side by side
    pair = doc.add_paragraph()
    pair.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pair.paragraph_format.space_before = Pt(8)
    pair.paragraph_format.space_after = Pt(2)
    pair.add_run().add_picture(str(print_copy("learn.png")), height=Cm(4.5))
    pair.add_run("    ")
    pair.add_run().add_picture(str(print_copy("analytics.png")), height=Cm(4.5))
    caption(doc, "The code lab, left: two of three tests pass on the first attempt, by design.   "
                 "Insights, right: completion by activity type and one recommended change.")

    heading(doc, "What happens next")
    para(doc, "Week 9: authentication, database schema, deployment and the generation pipeline, so a real "
              "course can be produced and saved.")
    para(doc, "Week 10: the code lab with real execution, the instructor editor, and analytics computed "
              "from real event data rather than seeded numbers.")
    para(doc, "Week 11: features freeze, public deployment, the five-minute demo video, the report, three "
              "user manuals and the slides.")

    foot = para(doc, "All screens are static and seeded with prototype data. Repository: "
                     "github.com/Rjiang12138/comp5241-project1  ·  Prototype: prototype/index.html",
                size=8, color=MUTED, before=10, after=0)
    set_border_bottom(foot, color="E4E7EC", size=4)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print("saved", OUT)


if __name__ == "__main__":
    build()
